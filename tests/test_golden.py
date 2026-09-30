"""Golden test: the engine reproduces every formula cell of the v0.4 workbook.

The workbook was last calculated with TODAY() = 2026-08-24 (its candidate intake date).
Evaluated as of that date, the engine must produce a value for every one of the
workbook's formula cells, on every sheet, equal to the result saved in the file.
"""

import collections
import datetime as dt
import unittest
import warnings
from pathlib import Path

from manifest_workbench import store as ST
from manifest_workbench.cells import workbook_cells
from manifest_workbench.engine import Workbench

SPEC = (Path(__file__).resolve().parent.parent / 'spec' /
        'Manifest_Workbench_CISC-001_MOPS-003_RCC-004_MASR_Registry_Candidate_Pipeline_v0_4.xlsx')
CALC_DATE = dt.date(2026, 8, 24)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def _norm(v):
    if v is None or v == '':
        return ''
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _equal(a, b):
    a, b = _norm(a), _norm(b)
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-9
    return a == b


@unittest.skipUnless(openpyxl and SPEC.exists(), 'openpyxl and the v0.4 workbook are required')
class GoldenWorkbookTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        from manifest_workbench.xlsx_import import import_workbook
        cls.store, cls.problems = import_workbook(SPEC)
        cls.cells = workbook_cells(Workbench(cls.store, CALC_DATE))
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            cls.formulas = openpyxl.load_workbook(SPEC)
            cls.values = openpyxl.load_workbook(SPEC, data_only=True)

    def test_import_is_clean(self):
        self.assertEqual(self.problems, [])
        counts = {k: len(self.store[k]) for k in ('portfolio', 'conviction', 'pew004', 'evidence',
                                                  'miar', 'review', 'masr', 'pipeline')}
        self.assertEqual(counts, {'portfolio': 47, 'conviction': 47, 'pew004': 20, 'evidence': 0,
                                  'miar': 47, 'review': 0, 'masr': 56, 'pipeline': 10})
        self.assertEqual(len(self.store['mandate']), 28)
        self.assertEqual(len([a for a in self.store['actions'] if a]), 6)

    def test_bundled_seed_matches_workbook(self):
        seed = ST.seed_store()
        for key in ST.TABLE_KEYS:
            self.assertEqual(seed[key], self.store[key], key)

    def test_every_formula_cell_on_every_sheet(self):
        checked, missing, mismatches = collections.Counter(), [], []
        for ws in self.formulas:
            ours = self.cells.get(ws.title, {})
            for row in ws.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith('=')):
                        continue
                    checked[ws.title] += 1
                    ref = f'{ws.title}!{cell.coordinate}'
                    if cell.coordinate not in ours:
                        missing.append(ref)
                        continue
                    expected = self.values[ws.title][cell.coordinate].value
                    if not _equal(ours[cell.coordinate], expected):
                        mismatches.append(f'{ref}: {ours[cell.coordinate]!r} != {expected!r}')
        self.assertEqual(sum(checked.values()), 13950)
        self.assertEqual(missing, [])
        self.assertEqual(mismatches, [])

    def test_every_sheet_is_covered(self):
        from manifest_workbench.cli import SHEETS
        self.assertEqual(sorted(s for s, *_ in SHEETS), sorted(self.formulas.sheetnames))


if __name__ == '__main__':
    unittest.main()
