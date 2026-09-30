"""Import operating inputs from a Manifest Workbench .xlsx file.

Only input cells (and the certified MFPDF schedule) are read; every formula cell
is ignored and recomputed by the engine. Requires ``openpyxl``.
"""

from __future__ import annotations

import datetime as dt
import warnings

from . import schema as S
from .engine import blank
from .store import empty_store

LINKED_MASR_FIELDS = set(S.MASR_CERTIFIED_LINKED)


def _clean(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        if value == '' or value.startswith('='):
            return None
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _read_row(ws, table: S.Table, row: int, skip=()) -> dict:
    out = {}
    for field in table.fields:
        if not field.stored or field.key in skip:
            continue
        value = _clean(ws[f'{field.col}{row}'].value)
        if value is not None:
            out[field.key] = value
    return out


def _check_choices(table: S.Table, record: dict, where: str, problems: list) -> None:
    for field in table.fields:
        value = record.get(field.key)
        if field.choices is None or value is None:
            continue
        if value not in field.choices:
            problems.append(f"{where}: {field.header} '{value}' is not a controlled value")


ORIGINAL = ('portfolio', 'conviction', 'miar', 'pew004', 'evidence', 'review', 'pipeline', 'masr')


def _import_operating_sheets(wb, store, problems, holdings, sheet):
    """CISC-001 and MOPS-002 operating inputs (every table added after the RCC chain)."""
    for name, table in S.TABLES.items():
        if name in ORIGINAL:
            continue
        ws = sheet(table)
        if table.mode == 'single':
            rec = {}
            for f in table.fields:
                if f.stored:
                    value = _clean(ws[f.col].value)
                    if value is not None:
                        rec[f.key] = value
            _check_choices(table, rec, table.sheet, problems)
            store[name] = rec
            continue
        count = holdings if table.mode == 'positional' else table.capacity
        rows = []
        for row in range(table.first_row, table.first_row + count):
            rec = _read_row(ws, table, row)
            _check_choices(table, rec, f'{table.sheet}!{row}', problems)
            rows.append(rec)
        store[name] = rows


DOC_SHEETS = ('05 Workbench Guide', '14 PEW Guide', '15 MOPS-002 Freeze Record', 'Executive Summary',
              'Rebalancing', 'Sources & Certification')


def extract_docs(path) -> dict:
    """Static reference sheets as rows of (column, text); formula cells keep their reference."""
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(path)
    docs = {}
    for name in DOC_SHEETS:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith('='):
                    cells.append([c.column_letter, {'cell': c.coordinate}])
                else:
                    cells.append([c.column_letter, _clean(v) if not isinstance(v, str) else v])
            if cells:
                rows.append([row[0].row, cells])
        docs[name] = rows
    return docs


def import_workbook(path) -> tuple:
    """Return (store, warnings) built from the workbook at ``path``."""
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - exercised only without the extra installed
        raise RuntimeError("Importing .xlsx files needs openpyxl: pip install openpyxl") from None
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(path)
    store = empty_store()
    problems = []

    def sheet(table):
        if table.sheet not in wb.sheetnames:
            raise ValueError(f"Worksheet '{table.sheet}' not found; is this a Manifest Workbench file?")
        return wb[table.sheet]

    # Positional tables: one row per certified holding.
    for table in (S.ALLOCATION, S.CONVICTION, S.MIAR):
        ws = sheet(table)
        for row in range(table.first_row, table.first_row + table.capacity):
            rec = _read_row(ws, table, row)
            _check_choices(table, rec, f'{table.sheet}!{row}', problems)
            store[table.name].append(rec)
    store['portfolio'] = [r for r in store['portfolio'] if not blank(r.get('symbol'))]
    n = len(store['portfolio'])
    store['conviction'] = store['conviction'][:n]
    store['miar'] = store['miar'][:n]

    ws = sheet(S.PEW004)
    for row in range(S.PEW004.first_row, S.PEW004.first_row + S.PEW004.capacity):
        rec = _read_row(ws, S.PEW004, row)
        rec.setdefault('candidate_id', f'CAND-{row - S.PEW004.first_row + 1:03d}')
        _check_choices(S.PEW004, rec, f'{S.PEW004.sheet}!{row}', problems)
        store['pew004'].append(rec)

    # Registers: one row per record, blank rows skipped.
    for table in (S.EVIDENCE, S.REVIEWS, S.PIPELINE):
        ws = sheet(table)
        for row in range(table.first_row, table.first_row + table.capacity):
            rec = _read_row(ws, table, row)
            if rec:
                _check_choices(table, rec, f'{table.sheet}!{row}', problems)
                store[table.name].append(rec)

    ws = sheet(S.MASR)
    for row in range(S.MASR.first_row, S.MASR.first_row + S.MASR.capacity):
        ticker_cell = ws[f'B{row}'].value
        linked = isinstance(ticker_cell, str) and ticker_cell.startswith('=')
        rec = _read_row(ws, S.MASR, row, skip=LINKED_MASR_FIELDS if linked else ())
        if linked:
            pos = row - S.MASR.first_row
            if pos >= n:
                continue
            rec = {'_linked': pos, **rec}
        elif not rec:
            continue
        _check_choices(S.MASR, rec, f'{S.MASR.sheet}!{row}', problems)
        store['masr'].append(rec)

    _import_operating_sheets(wb, store, problems, n, sheet)

    title = wb['22 RCC-004 Control Center']['A1'].value if '22 RCC-004 Control Center' in wb.sheetnames else None
    store['meta'] = {'source': str(path), 'title': title,
                     'imported': dt.date.today().isoformat()}
    return store, problems
